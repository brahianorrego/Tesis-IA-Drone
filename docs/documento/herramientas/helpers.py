"""Utilidades para intervenir el Word con colores (rojo, amarillo, verde y turquesa)."""
import copy
import re

import docx
from docx.enum.text import WD_COLOR_INDEX as WC
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from docx.text.run import Run

ROJO, AMARILLO, VERDE, TURQ = WC.RED, WC.YELLOW, WC.BRIGHT_GREEN, WC.TURQUOISE
LOG = []


class Doc:
    def __init__(self, path):
        self.d = docx.Document(path)
        self.P = list(self.d.paragraphs)

    def find(self, texto, desde=0, hasta=None, n=0):
        hits = [p for p in self.P[desde:hasta] if texto in p.text]
        if len(hits) <= n:
            raise KeyError(f"No encontrado: {texto!r}")
        return hits[n]


# ---------- runs ----------
def run_text(r):
    out = []
    for c in r:
        if c.tag == qn('w:t'):
            out.append(c.text or '')
        elif c.tag == qn('w:tab'):
            out.append('\t')
        elif c.tag in (qn('w:br'), qn('w:cr')):
            out.append('\n')
    return ''.join(out)


def _set_t(r, txt):
    t = r.find(qn('w:t'))
    t.text = txt
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')


def _splittable(r):
    ts = r.findall(qn('w:t'))
    otros = [c for c in r if c.tag not in (qn('w:rPr'), qn('w:t'))]
    return len(ts) == 1 and not otros


def _child_len(c):
    if c.tag == qn('w:t'):
        return len(c.text or '')
    if c.tag in (qn('w:tab'), qn('w:br'), qn('w:cr')):
        return 1
    return 0


def _split_run(r, off):
    """Parte el run `r` en la posición `off` de su texto; el resto pasa a un run nuevo."""
    rpr = r.find(qn('w:rPr'))
    der = OxmlElement('w:r')
    if rpr is not None:
        der.append(copy.deepcopy(rpr))
    pos = 0
    hijos = [c for c in r if c.tag != qn('w:rPr')]
    mover = []
    for c in hijos:
        L = _child_len(c)
        if mover:
            mover.append(c)
        elif pos + L <= off:
            pass
        elif c.tag == qn('w:t') and pos < off < pos + L:
            txt = c.text
            c.text = txt[:off - pos]
            c.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            t2 = OxmlElement('w:t')
            t2.text = txt[off - pos:]
            t2.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            der.append(t2)
            mover.append(None)
        else:
            mover.append(c)
        pos += L
    for c in mover:
        if c is not None:
            der.append(c)
    r.addnext(der)


def _split_at(p, k):
    pos = 0
    for r in p._p.findall(qn('w:r')):
        t = run_text(r)
        a, b = pos, pos + len(t)
        if a < k < b:
            _split_run(r, k - a)
            return
        pos = b


def hl(r_el, p, color):
    Run(r_el, p).font.highlight_color = color


def _clean_rpr(rpr):
    if rpr is None:
        return None
    rpr = copy.deepcopy(rpr)
    for tag in ('w:highlight', 'w:b', 'w:bCs', 'w:i', 'w:iCs', 'w:u', 'w:color', 'w:vertAlign', 'w:strike',
                'w:rStyle', 'w:shd'):
        for e in rpr.findall(qn(tag)):
            rpr.remove(e)
    return rpr


def _template_rpr(p):
    for r in p._p.iter(qn('w:r')):
        if r.find(qn('w:t')) is not None:
            return _clean_rpr(r.find(qn('w:rPr')))
    return None


def new_run(p, texto, color, bold=False, italic=False, rpr=None, size=None, after=None):
    r = OxmlElement('w:r')
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = OxmlElement('w:t')
    t.text = texto
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    if after is not None:
        after.addnext(r)
    else:
        p._p.append(r)
    run = Run(r, p)
    if color is not None:
        run.font.highlight_color = color
    if bold:
        run.font.bold = True
    if italic:
        run.font.italic = True
    if size:
        run.font.size = Pt(size)
    return r


REV = re.compile(r'(\[REVISAR:[^\]]*\])')


def add_text(p, texto, color, bold=False, italic=False, rpr=None, size=None, after=None):
    """Agrega texto; los [REVISAR: ...] salen en turquesa y negrita."""
    last = after
    for parte in REV.split(texto):
        if not parte:
            continue
        if parte.startswith('[REVISAR:'):
            last = new_run(p, parte, TURQ, bold=True, rpr=rpr, size=size, after=last)
        else:
            last = new_run(p, parte, color, bold=bold, italic=italic, rpr=rpr, size=size, after=last)
    return last


# ---------- párrafos ----------
def marcar(p, color=ROJO, nota_figura=True):
    tiene_texto = False
    for r in p._p.iter(qn('w:r')):
        hl(r, p, color)
        if run_text(r).strip():
            tiene_texto = True
    if color == ROJO and not tiene_texto and p._p.xpath('.//w:drawing') and nota_figura:
        new_run(p, ' [eliminar esta figura]', ROJO, bold=True, rpr=_template_rpr(p))


def _nuevo_p(anchor_el, ppr_de=None, antes=False):
    new_p = OxmlElement('w:p')
    if ppr_de is not None and ppr_de.pPr is not None:
        ppr = copy.deepcopy(ppr_de.pPr)
        for e in ppr.findall(qn('w:sectPr')):
            ppr.remove(e)
        new_p.append(ppr)
    if antes:
        anchor_el.addprevious(new_p)
    else:
        anchor_el.addnext(new_p)
    return new_p


def par_after(doc, anchor, texto='', color=AMARILLO, style=None, bold=False, italic=False, copiar_formato=True,
              size=None):
    """Párrafo nuevo después de `anchor` (Paragraph o elemento XML, p. ej. una tabla)."""
    if isinstance(anchor, Paragraph):
        el, parent, src = anchor._p, anchor._parent, (anchor._p if copiar_formato and style is None else None)
    else:
        el, parent, src = anchor, doc.d._body, None
    new_p = _nuevo_p(el, src)
    para = Paragraph(new_p, parent)
    if style:
        para.style = doc.d.styles[style]
    elif src is None:
        para.style = doc.d.styles['Párr.APA']
    rpr = _template_rpr(anchor) if isinstance(anchor, Paragraph) and style is None else None
    if texto:
        add_text(para, texto, color, bold=bold, italic=italic, rpr=rpr, size=size)
    return para


def par_before(doc, anchor, texto, color=AMARILLO, bold=False):
    new_p = _nuevo_p(anchor._p, None, antes=True)
    para = Paragraph(new_p, anchor._parent)
    para.style = doc.d.styles['Normal']
    add_text(para, texto, color, bold=bold)
    return para


def revisar(p, nota):
    add_text(p, f' [REVISAR: {nota}]', TURQ, bold=True, rpr=_template_rpr(p))


def reemplazar(doc, p, viejo, nuevo=None, n=0, color_nuevo=VERDE):
    """Pone `viejo` en rojo y `nuevo` en verde justo después. Si no se puede partir el texto,
    marca el párrafo completo en rojo y agrega el párrafo corregido en verde."""
    full = ''.join(run_text(r) for r in p._p.findall(qn('w:r')))
    idx = -1
    for _ in range(n + 1):
        idx = full.find(viejo, idx + 1)
    if idx < 0:
        raise KeyError(f'No está en el párrafo: {viejo!r} :: {p.text[:80]!r}')
    fin = idx + len(viejo)
    try:
        _split_at(p, fin)
        _split_at(p, idx)
    except ValueError:
        LOG.append(f'fallback: {viejo[:40]!r}')
        texto = p.text.replace(viejo, nuevo or '', 1)
        marcar(p)
        return par_after(doc, p, texto, VERDE)
    pos, ultimo = 0, None
    for r in p._p.findall(qn('w:r')):
        t = run_text(r)
        a, b = pos, pos + len(t)
        if a >= idx and b <= fin and t:
            hl(r, p, ROJO)
            ultimo = r
        pos = b
    if nuevo:
        rpr = _clean_rpr(ultimo.find(qn('w:rPr')))
        add_text(p, nuevo, color_nuevo, rpr=rpr, after=ultimo)
    return p


def insertar_tras(p, ancla, nuevo, color=AMARILLO):
    full = ''.join(run_text(r) for r in p._p.findall(qn('w:r')))
    idx = full.find(ancla)
    if idx < 0:
        raise KeyError(ancla)
    fin = idx + len(ancla)
    _split_at(p, fin)
    pos, ultimo = 0, None
    for r in p._p.findall(qn('w:r')):
        pos += len(run_text(r))
        if pos == fin:
            ultimo = r
            break
    add_text(p, nuevo, color, rpr=_clean_rpr(ultimo.find(qn('w:rPr'))), after=ultimo)


# ---------- tablas y figuras ----------
def _bordes(tbl):
    tblPr = tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{lado}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:color'), '000000')
        b.append(e)
    tblPr.append(b)


def tabla_after(doc, anchor_el, filas, anchos_cm, color=AMARILLO, size=9):
    t = doc.d.add_table(rows=len(filas), cols=len(filas[0]))
    t.style = doc.d.styles['Table Grid']
    # Bordes estilo APA (plantilla USB): solo líneas horizontales
    b = OxmlElement('w:tblBorders')
    for lado, val in (('top', 'single'), ('left', 'nil'), ('bottom', 'single'), ('right', 'nil'),
                      ('insideH', 'single'), ('insideV', 'nil')):
        e = OxmlElement(f'w:{lado}')
        e.set(qn('w:val'), val)
        if val != 'nil':
            e.set(qn('w:sz'), '4')
            e.set(qn('w:color'), '000000')
        b.append(e)
    look = t._tbl.tblPr.find(qn('w:tblLook'))
    if look is not None:
        look.addprevious(b)
    else:
        t._tbl.tblPr.append(b)
    for i, fila in enumerate(filas):
        for j, val in enumerate(fila):
            c = t.cell(i, j)
            c.width = Cm(anchos_cm[j])
            par = c.paragraphs[0]
            par.style = doc.d.styles['Normal']
            add_text(par, str(val), color, bold=(i == 0), size=size)
    t.autofit = False
    for j, w in enumerate(anchos_cm):
        t.columns[j].width = Cm(w)
    for fila in t.rows:
        for c in fila.cells:
            for par in c.paragraphs:
                par.alignment = 0
    anchor_el.addnext(t._tbl)
    return t._tbl


def figura_after(doc, anchor, img, ancho_cm=None, alto_max_cm=19.0, ancho_max_cm=16.0):
    from PIL import Image
    w, h = Image.open(img).size
    ancho = min(ancho_max_cm, alto_max_cm * w / h) if ancho_cm is None else ancho_cm
    new_p = _nuevo_p(anchor._p if isinstance(anchor, Paragraph) else anchor, None)
    para = Paragraph(new_p, doc.d._body)
    para.style = doc.d.styles['Normal']
    para.alignment = 1
    para.add_run().add_picture(img, width=Cm(ancho))
    return para


def titulo_fig(doc, anchor, etiqueta, titulo, color=AMARILLO):
    """Rótulo estilo APA: etiqueta en negrita y título en cursiva en la línea siguiente."""
    p1 = par_after(doc, anchor, etiqueta, color, bold=True, style='Normal')
    p2 = par_after(doc, p1, titulo, color, italic=True, style='Normal')
    return p2


def nota_after(doc, anchor, texto, color=AMARILLO):
    """Nota de tabla o figura según la plantilla USB: 'Nota.' en cursiva, todo en 10 puntos."""
    from docx.text.paragraph import Paragraph as _P
    el = anchor._p if isinstance(anchor, _P) else anchor
    new_p = _nuevo_p(el, None)
    para = Paragraph(new_p, doc.d._body)
    para.style = doc.d.styles['Normal']
    texto = texto[len('Nota.'):].lstrip() if texto.startswith('Nota.') else texto
    new_run(para, 'Nota.', color, italic=True, size=10)
    add_text(para, ' ' + texto, color, size=10)
    return para
